# assist/export_views.py
import re
import io
import os
import unicodedata

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, Http404
from django.shortcuts import get_object_or_404
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AssistMessage, AssistSession
from .markdown_blocks import parse_markdown_blocks, extract_tables

# ── AJi brand ───────────────────────────────────────────────────────────────
# Three colours, and only these three. No greys as fills, no accent palette:
# hierarchy comes from weight, size and space instead. Yellow marks the thing
# worth finding on a page — the masthead rule, a question, a table header.
AJI_BLACK = '000000'
AJI_YELLOW = 'FCD700'
AJI_WHITE = 'FFFFFF'

AJI_INK = AJI_BLACK  # the near-black it used to be is not a brand colour

LOGO_PATH = os.path.join(settings.BASE_DIR, 'static', 'img', 'logo.png')
LOGO_RATIO = 453 / 887  # the shipped logo, so a width is enough to place it

# Two faces, per the brand kit. Draft B is the display face and has no Arabic in
# it at all — 577 glyphs, not one of them Arabic script. IBM Plex Sans Arabic
# covers Arabic *and* Latin, so a line with any Arabic in it is set wholly in
# Plex rather than stitched together from two fonts mid-sentence.
FONT_BODY = 'DraftB-Regular'
FONT_BOLD = 'DraftB-Bold'
FONT_AR_BODY = 'PlexArabic-Regular'
FONT_AR_BOLD = 'PlexArabic-Bold'

# Word and Excel reference a font by name, so the family has to be installed on
# whichever machine opens the file. The PDF embeds them and always looks right.
DOCX_FONT = 'Draft B'
DOCX_FONT_AR = 'IBM Plex Sans Arabic'

_PDF_FONT_FILES = {
    FONT_BODY: 'DraftB-Regular.ttf',
    FONT_BOLD: 'DraftB-Bold.ttf',
    FONT_AR_BODY: 'PlexArabic-Regular.ttf',
    FONT_AR_BOLD: 'PlexArabic-Bold.ttf',
}

_PDF_FONTS_READY = False
_ARABIC_RE = re.compile(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]')


def _ensure_pdf_fonts():
    """Register the brand fonts with ReportLab, once.

    The brand kit ships these as .otf with PostScript (CFF) outlines, which
    ReportLab refuses outright; assist/build_brand_fonts.py converts them to
    TrueType and the .ttf files are committed under static/fonts/. A missing one
    is reported rather than quietly swapped for Helvetica — the brand kit is
    explicit that a silent system-font fallback is worse than a failure.
    """
    global _PDF_FONTS_READY
    if _PDF_FONTS_READY:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, filename in _PDF_FONT_FILES.items():
        path = os.path.join(settings.BASE_DIR, 'static', 'fonts', filename)
        if not os.path.exists(path):
            raise RuntimeError(
                f'AJi brand font missing: {filename}. Rebuild it with '
                f'"python -m assist.build_brand_fonts <brand-kit assets/fonts>".')
        pdfmetrics.registerFont(TTFont(name, path))
        _load_font_coverage(name, path)
    _PDF_FONTS_READY = True


def _fonts_for(text):
    """(regular, bold) for this text — Plex whenever it contains any Arabic."""
    if _ARABIC_RE.search(text or ''):
        return FONT_AR_BODY, FONT_AR_BOLD
    return FONT_BODY, FONT_BOLD


_FONT_CMAPS = {}

# Common "look-alike" characters an LLM might emit that aren't in the bundled
# font — map them to a safe equivalent so nothing ever renders as a black box.
_GLYPH_FALLBACKS = {
    '\u3010': '[', '\u3011': ']',   # 【 】 CJK brackets
    '\uFF08': '(', '\uFF09': ')',   # full-width parens
    '\u2016': '|', '\u2225': '|',   # double vertical bar
    '\u2018': "'", '\u2019': "'",   # curly single quotes
    '\u201C': '"', '\u201D': '"',   # curly double quotes
    '\u2026': '...',                # ellipsis
    '\u2022': '-',                  # bullet
    '\u2013': '-', '\u2014': '-',   # en/em dash (kept as '-' if font lacks them)
}


def _load_font_coverage(font_name, font_path):
    try:
        from fontTools.ttLib import TTFont as FTFont
        with FTFont(font_path, lazy=True) as f:
            _FONT_CMAPS[font_name] = set(f.getBestCmap().keys())
    except Exception:
        # Unknown coverage — skip sanitisation for this font rather than fail.
        _FONT_CMAPS[font_name] = None


def _sanitize_for_font(text, font_name):
    """Replace any character this font can't draw with a safe fallback, so no
    glyph ever renders as a missing-character box.

    Which font matters now that there are two: Draft B has no Arabic, so the
    coverage question has a different answer depending on what is being set.
    """
    cmap = _FONT_CMAPS.get(font_name)
    if cmap is None or not text:
        return text
    out = []
    for ch in text:
        if ord(ch) in cmap or ch in ('\n', '\t'):
            out.append(ch)
            continue
        mapped = _GLYPH_FALLBACKS.get(ch)
        if mapped is None:
            # Anything the table misses, answer by what kind of character it is.
            # A knowledge base writes "Part‑G" with a non-breaking hyphen and
            # "0.75 m" with a narrow no-break space; both were reaching the page
            # as "?" because a display face carries neither.
            category = unicodedata.category(ch)
            mapped = (' ' if category == 'Zs'
                      else '-' if category == 'Pd'
                      else '' if category in ('Cf', 'Mn')
                      else '?')
        out.append(mapped)
    return ''.join(out)


def _rtl_shape(text):
    """Reshape + bidi-reorder Arabic text so it renders as connected,
    right-to-left script instead of disjointed left-to-right letterforms.
    Non-Arabic text is returned unchanged."""
    if not text or not _ARABIC_RE.search(text):
        return text
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(text))


def _plain_text(md_text):
    """Strip the light markdown AJiAssist uses (bold/italic/headers/links) for
    plain-text rendering in docx/pdf paragraphs."""
    t = md_text or ''
    t = re.sub(r'```.*?```', '', t, flags=re.S)
    # The knowledge base writes multi-line table cells as HTML line breaks.
    # Nothing turned them back into breaks, so "<br>" was printing as itself.
    t = re.sub(r'<br\s*/?>', '\n', t, flags=re.I)
    t = re.sub(r'^#{1,6}\s*', '', t, flags=re.M)
    t = re.sub(r'\*\*(.+?)\*\*', r'\1', t)
    t = re.sub(r'\*(.+?)\*', r'\1', t)
    t = re.sub(r'~~(.+?)~~', r'\1', t)
    t = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', t)
    return t.strip()


def _safe_filename(name, ext):
    name = re.sub(r'[^\w\- ]+', '', name or 'AJiAssist').strip() or 'AJiAssist'
    return f'{name[:60]}.{ext}'


# ── Single-answer table → Excel ────────────────────────────────────────────

@login_required
def export_message_table_xlsx(request, msg_pk):
    msg = get_object_or_404(AssistMessage, pk=msg_pk, user=request.user)
    tables = extract_tables(msg.answer)
    if not tables:
        raise Http404('This answer has no table to export.')

    wb = Workbook()
    wb.properties.creator = 'AJi Assist'
    wb.remove(wb.active)
    for idx, rows in enumerate(tables, start=1):
        ws = wb.create_sheet(title=f'Table {idx}'[:31])
        header, *body = rows
        # This one leaves on its own, so it carries the masthead. The table
        # sheets inside a conversation workbook do not — sheet one already has.
        start = _xl_masthead(
            ws, _plain_text(msg.question),
            f'AJi Assist · exported {timezone.now():%d %B %Y, %H:%M}', len(header))
        _write_branded_table(ws, header, body, start_row=start)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{_safe_filename(msg.question, "xlsx")}"'
    return resp


# ── Excel: the brand, applied ───────────────────────────────────────────────

def _xl_font(size=10, bold=False, white=False, arabic=False):
    return Font(name=DOCX_FONT_AR if arabic else DOCX_FONT, size=size, bold=bold,
                color=AJI_WHITE if white else AJI_BLACK)


def _xl_border():
    from openpyxl.styles import Border, Side
    thin = Side(style='thin', color=AJI_BLACK)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _xl_header_row(ws, row, ncols):
    """Yellow band, black bold text. Excel's own table styles are blue."""
    from openpyxl.styles import Alignment
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = _xl_font(10, bold=True)
        cell.fill = PatternFill('solid', fgColor=AJI_YELLOW)
        cell.border = _xl_border()
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 22


def _xl_masthead(ws, title, subtitle, ncols):
    """Logo, title, and a yellow rule — the same opening as the other two
    formats. Returns the row the content should start on."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter as gcl

    if os.path.exists(LOGO_PATH):
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(LOGO_PATH)
            img.width = 132
            img.height = int(132 * LOGO_RATIO)
            ws.add_image(img, 'A1')
        except Exception:
            pass  # a workbook without the mark still beats no workbook
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 20

    last = gcl(max(ncols, 1))
    ws.merge_cells(f'A5:{last}5')
    ws['A5'] = title
    ws['A5'].font = _xl_font(16, bold=True)
    ws['A5'].alignment = Alignment(vertical='center')
    ws.row_dimensions[5].height = 24

    ws.merge_cells(f'A6:{last}6')
    ws['A6'] = subtitle
    ws['A6'].font = _xl_font(8.5)

    # The rule: a short row filled yellow all the way across.
    for col in range(1, ncols + 1):
        ws.cell(row=7, column=col).fill = PatternFill('solid', fgColor=AJI_YELLOW)
    ws.row_dimensions[7].height = 4
    return 8


def _write_branded_table(ws, header, body, start_row=1):
    """A header band and a bordered body, sized to its content."""
    from openpyxl.styles import Alignment
    for col, value in enumerate(header, start=1):
        ws.cell(row=start_row, column=col, value=_plain_text(value))
    _xl_header_row(ws, start_row, len(header))

    for r, row in enumerate(body, start=start_row + 1):
        for col in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=col,
                           value=_plain_text(row[col - 1]) if col - 1 < len(row) else '')
            cell.font = _xl_font(10)
            cell.border = _xl_border()
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col in range(1, len(header) + 1):
        widest = max(
            [len(str(ws.cell(row=r, column=col).value or ''))
             for r in range(start_row, start_row + len(body) + 1)] or [10])
        ws.column_dimensions[get_column_letter(col)].width = min(60, max(12, widest + 2))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


# ── Word: the bits python-docx has no API for ──────────────────────────────

def _w(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _shade(properties, fill):
    """Paint a paragraph or a table cell. Both take the same w:shd element,
    they just hang it off different property nodes."""
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(_w('w:val'), 'clear')
    shd.set(_w('w:color'), 'auto')
    shd.set(_w('w:fill'), fill)
    properties.append(shd)


def _rule_below(paragraph, fill, eighths=24):
    """A coloured rule under a paragraph — the masthead bar. Word measures
    border weight in eighths of a point, so 24 is a 3pt bar."""
    from docx.oxml import OxmlElement
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(_w('w:val'), 'single')
    bottom.set(_w('w:sz'), str(eighths))
    bottom.set(_w('w:space'), '4')
    bottom.set(_w('w:color'), fill)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _page_number_field(paragraph):
    """Word writes page numbers as a field, which python-docx cannot add."""
    from docx.oxml import OxmlElement
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(_w('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(_w('xml:space'), 'preserve')
    instr.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(_w('w:fldCharType'), 'end')
    for node in (begin, instr, end):
        run._r.append(node)
    return run


def _use_brand_font(run, arabic=False):
    """Name the face on every script slot, or Word quietly sets Latin in Draft B
    and everything else in whatever its theme says."""
    name = DOCX_FONT_AR if arabic else DOCX_FONT
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(_w('w:rFonts'))
    if fonts is None:
        from docx.oxml import OxmlElement
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for slot in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(_w(slot), name)


def _use_brand_font_style(style):
    """The same every-script naming, for a style rather than a run."""
    from docx.oxml import OxmlElement
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(_w('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for slot, name in (('w:ascii', DOCX_FONT), ('w:hAnsi', DOCX_FONT),
                       ('w:cs', DOCX_FONT_AR), ('w:eastAsia', DOCX_FONT)):
        fonts.set(_w(slot), name)


def _brand_run(paragraph, text, size, bold=False, white=False):
    from docx.shared import Pt, RGBColor
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if white else RGBColor(0, 0, 0)
    _use_brand_font(run, arabic=bool(_ARABIC_RE.search(text or '')))
    return run


def _cell_text(paragraph, text, size, bold=False):
    """A cell's content, with its line breaks kept. Word needs an explicit break
    element — a newline inside a run is just whitespace to it.

    The break goes in a run of its own. Putting it inside the text run means
    inserting it ahead of that run's w:rPr, and a run whose properties do not
    come first is invalid enough that Word offers to repair the file.
    """
    for i, line in enumerate((text or '').split('\n')):
        if i:
            _brand_run(paragraph, '', size, bold=bold).add_break()
        _brand_run(paragraph, line, size, bold=bold)
    return paragraph


def _docx_masthead(doc, title, subtitle):
    """Logo, a yellow bar, then the title. The same masthead the PDF opens with."""
    from docx.shared import Inches, Pt

    logo_p = doc.add_paragraph()
    logo_p.paragraph_format.space_after = Pt(4)
    if os.path.exists(LOGO_PATH):
        logo_p.add_run().add_picture(LOGO_PATH, width=Inches(1.45))
    _rule_below(logo_p, AJI_YELLOW)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(14)
    title_p.paragraph_format.space_after = Pt(2)
    _brand_run(title_p, title, 21, bold=True)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    _brand_run(sub_p, subtitle, 8.5)


# ── Full conversation exports ──────────────────────────────────────────────

def _session_messages(session_pk, user):
    session = get_object_or_404(AssistSession, pk=session_pk, user=user)
    messages = session.messages.filter(answer__gt='').order_by('created_at')
    return session, messages


@login_required
def export_session_docx(request, session_pk):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    session, messages = _session_messages(session_pk, request.user)
    title = session.name or session.auto_name()

    doc = Document()

    # Draft B as the document's own default, so anything added later inherits it
    # rather than Word's Calibri.
    normal = doc.styles['Normal']
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(10.5)
    _use_brand_font_style(normal)

    _docx_masthead(doc, title,
                   f'AJi Assist conversation · exported {timezone.now():%d %B %Y, %H:%M}')

    for msg in messages:
        # The question, in a yellow band. It is the thing a reader scans for
        # when scrolling back through a long conversation.
        qp = doc.add_paragraph()
        fmt = qp.paragraph_format
        fmt.space_before = Pt(14)
        fmt.space_after = Pt(8)
        fmt.left_indent = Pt(8)
        fmt.right_indent = Pt(8)
        _shade(qp._p.get_or_add_pPr(), AJI_YELLOW)
        _brand_run(qp, msg.question, 11.5, bold=True)

        for block in parse_markdown_blocks(msg.answer):
            if block['type'] == 'text':
                if block['content']:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(8)
                    _brand_run(p, _plain_text(block['content']), 10.5)
            else:
                rows = block['rows']
                header, *body = rows
                ncols = len(header)
                table = doc.add_table(rows=1, cols=ncols)
                # Table Grid, not one of Word's blue "Accent" styles: those
                # bring their own theme colour, which is not one of the three.
                table.style = 'Table Grid'
                for i, cell_text in enumerate(header):
                    cell = table.rows[0].cells[i]
                    cell.text = ''
                    _shade(cell._tc.get_or_add_tcPr(), AJI_YELLOW)
                    _cell_text(cell.paragraphs[0], _plain_text(cell_text), 9.5, bold=True)
                for row in body:
                    cells = table.add_row().cells
                    for i in range(ncols):
                        cells[i].text = ''
                        _cell_text(cells[i].paragraphs[0],
                                   _plain_text(row[i] if i < len(row) else ''), 9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Footer: where it came from, and where you are in it.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _brand_run(footer, 'AJi Assist · AJi Group · page ', 8)
    _use_brand_font(_page_number_field(footer))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    resp['Content-Disposition'] = f'attachment; filename="{_safe_filename(title, "docx")}"'
    return resp


@login_required
def export_session_pdf(request, session_pk):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image, KeepTogether)
    from reportlab.lib.styles import ParagraphStyle

    _ensure_pdf_fonts()
    session, messages = _session_messages(session_pk, request.user)
    title = session.name or session.auto_name()

    yellow = colors.HexColor('#FCD700')
    black = colors.HexColor('#000000')

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.6*cm, bottomMargin=1.8*cm,
                            leftMargin=2*cm, rightMargin=2*cm,
                            title=title, author='AJi Group')
    frame_w = doc.width

    def esc(t):
        return (t or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def prepared(t, bold=False):
        """Pick the face this text needs, drop anything that face cannot draw,
        shape Arabic, then escape for ReportLab's mini-XML. Line breaks survive
        as the one tag left standing after escaping."""
        regular, strong = _fonts_for(t)
        font = strong if bold else regular
        text = esc(_rtl_shape(_sanitize_for_font(t or '', font)))
        return text.replace('\n', '<br/>'), font

    def para(t, size, bold=False, leading=None, space_after=6, colour=black,
             back=None, pad=0):
        text, font = prepared(t, bold)
        style = ParagraphStyle(
            f'p{size}{bold}{back}', fontName=font, fontSize=size,
            leading=leading or size * 1.45, textColor=colour,
            spaceAfter=space_after, backColor=back,
            borderPadding=pad, wordWrap='RTL' if _ARABIC_RE.search(t or '') else None)
        return Paragraph(text, style)

    def masthead():
        """Logo over a yellow bar — the same opening as the Word document."""
        bits = []
        if os.path.exists(LOGO_PATH):
            w = 3.6 * cm
            bits.append(Image(LOGO_PATH, width=w, height=w * LOGO_RATIO))
        bar = Table([['']], colWidths=[frame_w], rowHeights=[3])
        bar.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), yellow),
                                 ('LINEBELOW', (0, 0), (-1, -1), 0, yellow),
                                 ('TOPPADDING', (0, 0), (-1, -1), 0),
                                 ('BOTTOMPADDING', (0, 0), (-1, -1), 0)]))
        bits += [Spacer(1, 6), bar, Spacer(1, 14)]
        return bits

    def furniture(canvas, _doc):
        """Every page carries the mark and its number. Black and yellow only."""
        canvas.saveState()
        canvas.setFillColor(black)
        canvas.setFont(FONT_BODY, 7.5)
        canvas.drawString(doc.leftMargin, 1.05 * cm, 'AJi Assist · AJi Group')
        canvas.drawRightString(A4[0] - doc.rightMargin, 1.05 * cm, str(canvas.getPageNumber()))
        canvas.setStrokeColor(yellow)
        canvas.setLineWidth(1.6)
        canvas.line(doc.leftMargin, 1.45 * cm, A4[0] - doc.rightMargin, 1.45 * cm)
        canvas.restoreState()

    story = masthead()
    story.append(para(title, 20, bold=True, space_after=2))
    story.append(para(f'AJi Assist conversation · exported {timezone.now():%d %B %Y, %H:%M}',
                      8.5, space_after=16))

    for msg in messages:
        # The question in a yellow band, the width of the frame.
        band = Table([[para(msg.question, 11, bold=True, space_after=0)]],
                     colWidths=[frame_w])
        band.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), yellow),
            ('LEFTPADDING', (0, 0), (-1, -1), 9),
            ('RIGHTPADDING', (0, 0), (-1, -1), 9),
            ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story += [Spacer(1, 10), band, Spacer(1, 9)]

        for block in parse_markdown_blocks(msg.answer):
            if block['type'] == 'text':
                for chunk in block['content'].split('\n\n'):
                    chunk = chunk.strip()
                    if not chunk:
                        continue
                    lines = [prepared(_plain_text(line))[0] for line in chunk.split('\n')]
                    joined = '<br/>'.join(lines)
                    _, font = prepared(chunk)
                    story.append(Paragraph(joined, ParagraphStyle(
                        'body', fontName=font, fontSize=10, leading=14.5,
                        textColor=black, spaceAfter=7)))
            else:
                rows = block['rows']
                # Cells carry markdown too — "**General Design**" was reaching
                # the page with its asterisks on.
                data = [[para(_plain_text(c), 8.5, bold=(ri == 0), space_after=0)
                         for c in r] for ri, r in enumerate(rows)]
                t = Table(data, hAlign='LEFT', repeatRows=1)
                t.setStyle(TableStyle([
                    # Yellow header with black text, black grid, white body —
                    # the three colours doing all of the work.
                    ('BACKGROUND', (0, 0), (-1, 0), yellow),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                    ('GRID', (0, 0), (-1, -1), 0.6, black),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ]))
                story += [Spacer(1, 3), t, Spacer(1, 10)]

    doc.build(story, onFirstPage=furniture, onLaterPages=furniture)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{_safe_filename(title, "pdf")}"'
    return resp


@login_required
def export_session_xlsx(request, session_pk):
    session, messages = _session_messages(session_pk, request.user)

    from openpyxl.styles import Alignment

    title = session.name or session.auto_name()
    wb = Workbook()
    wb.properties.creator = 'AJi Assist'
    ws = wb.active
    ws.title = 'Conversation'

    headings = ['#', 'Asked', 'Question', 'Answer']
    head_row = _xl_masthead(
        ws, title,
        f'AJi Assist conversation · exported {timezone.now():%d %B %Y, %H:%M}',
        len(headings))

    for col, value in enumerate(headings, start=1):
        ws.cell(row=head_row, column=col, value=value)
    _xl_header_row(ws, head_row, len(headings))

    table_sheets = []
    row = head_row + 1
    for n, msg in enumerate(messages, start=1):
        blocks = parse_markdown_blocks(msg.answer)
        text_parts = [b['content'] for b in blocks if b['type'] == 'text']
        values = [n, msg.created_at.strftime('%d %b %Y %H:%M'), msg.question,
                  _plain_text('\n\n'.join(text_parts))]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = _xl_font(10, arabic=bool(_ARABIC_RE.search(str(value))))
            cell.border = _xl_border()
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        row += 1
        for block in blocks:
            if block['type'] == 'table':
                table_sheets.append((n, block['rows']))

    for column, width in (('A', 5), ('B', 17), ('C', 42), ('D', 74)):
        ws.column_dimensions[column].width = width
    # The headings stay put while a long conversation scrolls under them.
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    # A table inside an answer is worth its own sheet — it is the thing somebody
    # opened Excel for.
    for n, rows in table_sheets:
        ws2 = wb.create_sheet(title=f'Table (Q{n})'[:31])
        header, *body = rows
        _write_branded_table(ws2, header, body, start_row=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{_safe_filename(title, "xlsx")}"'
    return resp
