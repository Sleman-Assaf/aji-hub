# people/management/commands/import_birthdays.py
"""Fill in employee birthdays, and nothing else.

    manage.py import_birthdays --dry-run
    manage.py import_birthdays

Every one of the 1,256 employee records has an empty `birthday`, so the
Birthdays block on the homepage has never had anything to show — not "nobody
today", but nobody ever. The cause is which file was imported: `import_employees`
defaults to the sheet "Consolidated-8 entities" in aji_hr_data.xlsx, and that
sheet has no birth date column at all. The dates are in the other workbook,
AJi_Combined_Data.xlsx, on a sheet with different columns in a different order.

Re-running `import_employees` against that workbook is not the fix. It writes
names, departments, sections, managers and user accounts from column positions
that do not line up with this sheet, so it would repair the birthdays by
damaging everything around them.

This reads one column, matches on employee code, and writes one field. An
employee it cannot match, or a date it cannot read, is left exactly as it was.
"""
from datetime import datetime

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError

from people.models import Employee

DEFAULT_SHEET = 'Employees'
CODE_COLUMN = 'Employee Code'
BIRTH_COLUMN = 'Birth Date'


def parse_birthday(value):
    """"1989-01-11" or a datetime to "01-11". Anything else to ''.

    Only the month and day are kept. The year is somebody's age, which is not
    the Hub's business and not what the block needs.
    """
    if not value:
        return ''
    try:
        if isinstance(value, datetime):
            return value.strftime('%m-%d')
        return datetime.strptime(str(value)[:10], '%Y-%m-%d').strftime('%m-%d')
    except (ValueError, TypeError):
        return ''


class Command(BaseCommand):
    help = "Fills in Employee.birthday from the combined data workbook. Touches no other field."

    def add_arguments(self, parser):
        parser.add_argument(
            '--file', '-f',
            default=str(settings.BASE_DIR / 'AJi_Combined_Data.xlsx'),
            help='Workbook to read (default: AJi_Combined_Data.xlsx beside manage.py).')
        parser.add_argument('--sheet', default=DEFAULT_SHEET,
                            help=f'Sheet to read (default: "{DEFAULT_SHEET}").')
        parser.add_argument('--dry-run', action='store_true',
                            help='Report what would change and write nothing.')

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError('openpyxl is needed to read the workbook.')

        path, sheet, dry = options['file'], options['sheet'], options['dry_run']
        try:
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except OSError as exc:
            raise CommandError(f'Could not open {path}: {exc}')
        if sheet not in book.sheetnames:
            raise CommandError(
                f'No sheet called "{sheet}" in {path}. Found: {", ".join(book.sheetnames)}')

        rows = book[sheet].iter_rows(values_only=True)
        header = list(next(rows, ()) or ())
        for column in (CODE_COLUMN, BIRTH_COLUMN):
            if column not in header:
                raise CommandError(
                    f'Sheet "{sheet}" has no "{column}" column. Found: '
                    f'{", ".join(str(h) for h in header if h)}')
        code_at, birth_at = header.index(CODE_COLUMN), header.index(BIRTH_COLUMN)

        # Read the sheet first, then go to the database once. Matching row by row
        # would be 1,200 queries to set 1,200 fields.
        wanted = {}
        unreadable = 0
        for row in rows:
            code = row[code_at]
            if code in (None, ''):
                continue
            day = parse_birthday(row[birth_at])
            if day:
                wanted[str(code).strip()] = day
            else:
                unreadable += 1

        employees = {e.emp_code.strip(): e for e in
                     Employee.objects.exclude(emp_code='').exclude(emp_code__isnull=True)
                     if e.emp_code}

        changed, already, unmatched = [], 0, []
        for code, day in wanted.items():
            person = employees.get(code)
            if person is None:
                unmatched.append(code)
            elif person.birthday == day:
                already += 1
            else:
                person.birthday = day
                changed.append(person)

        w = self.stdout.write
        w(f'  read from sheet   {len(wanted)}')
        w(f'  already correct   {already}')
        w(f'  to set            {len(changed)}')
        if unreadable:
            w(f'  unreadable dates  {unreadable} (left alone)')
        if unmatched:
            w(f'  no such employee  {len(unmatched)} '
              f'({", ".join(unmatched[:5])}{"…" if len(unmatched) > 5 else ""})')

        if dry:
            w(self.style.WARNING('Dry run — nothing was written.'))
            return
        if not changed:
            w(self.style.SUCCESS('Nothing to do.'))
            return

        Employee.objects.bulk_update(changed, ['birthday'], batch_size=500)
        w(self.style.SUCCESS(f'{len(changed)} birthdays filled in.'))
