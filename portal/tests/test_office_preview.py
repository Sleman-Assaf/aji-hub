"""Previewing a spreadsheet or a Word document without downloading it.

The point of the feature is checking a file before you save it, so these assert
the two things that would make it useless: that the values actually appear, and
that a document nobody may download still previews. The escaping test is the one
that matters most — a cell is text somebody typed into a file we unzip.
"""
import io
import logging

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from portal.models import Policy, PolicyCategory


def setUpModule():
    logging.disable(logging.CRITICAL)


def tearDownModule():
    logging.disable(logging.NOTSET)


def a_workbook(sheets):
    """{name: [[row], [row]]} as a real .xlsx in memory."""
    import openpyxl
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    book.save(buf)
    buf.seek(0)
    return SimpleUploadedFile(
        'book.xlsx', buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def a_word_file(blocks):
    """A real .docx. A tuple ('h', text) is a heading, ('t', rows) a table."""
    import docx
    document = docx.Document()
    for block in blocks:
        if isinstance(block, tuple) and block[0] == 'h':
            document.add_heading(block[1], level=1)
        elif isinstance(block, tuple) and block[0] == 't':
            rows = block[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(block)
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return SimpleUploadedFile(
        'notes.docx', buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


class OfficePreviewTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.group = Group.objects.create(name='Policies')
        cls.sana = User.objects.create_user('sana', password='x')
        cls.sana.groups.add(cls.group)
        cls.reader = User.objects.create_user('reader', password='x')

    def setUp(self):
        self.client.force_login(self.reader)
        self.lib = PolicyCategory.objects.create(name='Manuals',
                                                 created_by=self.sana)

    def a_doc(self, upload, **kwargs):
        return Policy.objects.create(category=self.lib, title='Thing',
                                     file=upload, **kwargs)

    # ── what the viewer is told ─────────────────────────────────────────────

    def test_a_workbook_reports_one_page_per_sheet(self):
        doc = self.a_doc(a_workbook({'Summary': [['a']],
                                     'Detail': [['b']],
                                     'Notes': [['c']]}))
        meta = self.client.get(reverse('policy_document', args=[doc.pk])).json()
        self.assertEqual(meta['kind'], 'sheet')
        self.assertEqual(meta['pages'], 3)
        self.assertEqual(meta['labels'], ['Summary', 'Detail', 'Notes'])
        self.assertTrue(meta['as_html'])

    def test_a_word_file_is_one_page(self):
        doc = self.a_doc(a_word_file(['Hello there.']))
        meta = self.client.get(reverse('policy_document', args=[doc.pk])).json()
        self.assertEqual(meta['kind'], 'doc')
        self.assertEqual(meta['pages'], 1)
        self.assertTrue(meta['as_html'])

    def test_a_format_that_cannot_be_shown_says_so(self):
        """.doc and .xls are the old binary formats and have no reader here.

        Claiming them would open an empty viewer instead of offering the file.
        """
        for name in ('old.doc', 'old.xls', 'bundle.zip'):
            doc = self.a_doc(SimpleUploadedFile(name, b'not really'))
            meta = self.client.get(reverse('policy_document', args=[doc.pk])).json()
            self.assertEqual(meta['kind'], '', name)
            self.assertEqual(meta['pages'], 0, name)
            doc.delete()

    # ── what comes back ─────────────────────────────────────────────────────

    def test_a_sheet_renders_its_values(self):
        doc = self.a_doc(a_workbook({'Costs': [['Item', 'Amount'],
                                               ['Cement', 1200],
                                               ['Steel', 34500]]}))
        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertIn('Costs', html)                    # the sheet's name
        self.assertIn('Cement', html)
        self.assertIn('1,200', html)                    # grouped, not 1200.0
        self.assertIn('34,500', html)
        self.assertIn('>A<', html)                      # column letters
        self.assertIn('ofc-rownum', html)               # and row numbers

    def test_each_sheet_is_its_own_page(self):
        doc = self.a_doc(a_workbook({'One': [['first sheet']],
                                     'Two': [['second sheet']]}))
        first = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        second = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 1])).content.decode()
        self.assertIn('first sheet', first)
        self.assertNotIn('second sheet', first)
        self.assertIn('second sheet', second)

    def test_a_sheet_beyond_the_last_one_is_not_found(self):
        doc = self.a_doc(a_workbook({'Only': [['x']]}))
        res = self.client.get(reverse('policy_doc_page', args=[doc.pk, 4]))
        self.assertEqual(res.status_code, 404)

    def test_a_word_file_keeps_headings_and_table_order(self):
        """Tables in the right place, not all at the end.

        python-docx exposes paragraphs and tables as two separate lists, so
        building from those puts every table after all the text — a form whose
        fields sat between its instructions came out as instructions then fields.
        """
        doc = self.a_doc(a_word_file([
            'Before the table.',
            ('h', 'Section One'),
            ('t', [['Name', 'Role'], ['Sana', 'Engineer']]),
            'After the table.',
        ]))
        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertIn('Section One', html)
        self.assertIn('Sana', html)
        self.assertLess(html.index('Before the table'), html.index('<table'))
        self.assertLess(html.index('<table'), html.index('After the table'))
        self.assertIn('<h3>Section One</h3>', html)

    def test_a_merged_cell_is_rendered_once(self):
        """python-docx returns a merged cell once per column it covers.

        A heading spanning three columns came out as the same words printed
        three times, so a job description with two merged columns rendered as
        five with everything duplicated.
        """
        import docx
        document = docx.Document()
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = 'Spans the lot'
        for c, value in enumerate(['one', 'two', 'three']):
            table.cell(1, c).text = value
        buf = io.BytesIO()
        document.save(buf)
        buf.seek(0)
        doc = self.a_doc(SimpleUploadedFile('merged.docx', buf.read()))

        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertEqual(html.count('Spans the lot'), 1)
        self.assertIn('colspan="3"', html)
        for value in ('one', 'two', 'three'):
            self.assertEqual(html.count('>' + value + '<'), 1, value)

    # ── the part that must not go wrong ─────────────────────────────────────

    def test_a_cell_cannot_carry_markup_into_the_page(self):
        """A cell is text somebody typed into a file we unzip and render."""
        nasty = '<img src=x onerror="alert(1)">'
        doc = self.a_doc(a_workbook({'Sheet': [['ok'], [nasty]]}))
        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertNotIn('<img src=x', html)
        self.assertIn('&lt;img src=x', html)
        self.assertNotIn('onerror="alert', html)

    def test_a_paragraph_cannot_carry_markup_either(self):
        doc = self.a_doc(a_word_file(['<script>alert(1)</script>']))
        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertNotIn('<script>alert', html)
        self.assertIn('&lt;script&gt;', html)

    def test_a_read_only_document_still_previews(self):
        """The whole point: check it before deciding whether to ask for it.

        A view-only document sends its rendered markup and never its bytes, the
        same bargain a view-only PDF already made.
        """
        self.lib.downloadable = False
        self.lib.save(update_fields=['downloadable'])
        doc = self.a_doc(a_workbook({'Sheet': [['secret value']]}))

        meta = self.client.get(reverse('policy_document', args=[doc.pk])).json()
        self.assertFalse(meta['may_download'])
        self.assertEqual(meta['download_url'], '')
        self.assertEqual(meta['pages'], 1)

        html = self.client.get(
            reverse('policy_doc_page', args=[doc.pk, 0])).content.decode()
        self.assertIn('secret value', html)
        # And the file itself is still refused.
        self.assertEqual(
            self.client.get(reverse('policy_doc_file', args=[doc.pk])).status_code,
            404)

    def test_a_document_nobody_may_read_previews_for_nobody(self):
        from portal.models import DocAccess
        self.lib.access = DocAccess.RESTRICTED
        self.lib.save(update_fields=['access'])
        doc = self.a_doc(a_workbook({'Sheet': [['secret value']]}))
        for name in ('policy_document',):
            self.assertEqual(
                self.client.get(reverse(name, args=[doc.pk])).status_code, 404)
        res = self.client.get(reverse('policy_doc_page', args=[doc.pk, 0]))
        self.assertEqual(res.status_code, 404)
